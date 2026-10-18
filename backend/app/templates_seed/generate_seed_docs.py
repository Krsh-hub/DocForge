"""Generator script to create seed DOCX templates for Mutual NDA, One-Way NDA, and SOW"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

seed_dir = Path(__file__).parent
seed_dir.mkdir(parents=True, exist_ok=True)

def style_heading(p, text, size=18, bold=True, color=(30, 41, 59)):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(*color)
    return run

def add_legal_disclaimer(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DRAFT TEMPLATE: REVIEW BY LEGAL COUNSEL BEFORE USE")
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(180, 83, 9) # Amber
    doc.add_paragraph() # Spacer

# =========================================================================
# 1. MUTUAL NDA
# =========================================================================
doc_mnda = Document()
add_legal_disclaimer(doc_mnda)

h1 = doc_mnda.add_paragraph()
style_heading(h1, "MUTUAL NON-DISCLOSURE AGREEMENT", size=18, bold=True)

sub = doc_mnda.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = sub.add_run("Reference: {{ doc_number }}")
r_sub.font.size = Pt(9.5)
r_sub.font.color.rgb = RGBColor(100, 116, 139)

doc_mnda.add_paragraph()

p_intro = doc_mnda.add_paragraph()
p_intro.add_run("This Mutual Non-Disclosure Agreement (\"Agreement\") is entered into on ")
p_intro.add_run("{{ effective_date }}").bold = True
p_intro.add_run(" (\"Effective Date\"), by and between:\n\n")

p_intro.add_run("1. {{ company.legal_name }}").bold = True
p_intro.add_run(", having its principal office at {{ company.address }}, {{ company.state }}, India (hereinafter referred to as \"Company\"); and\n\n")

p_intro.add_run("2. {{ client.legal_name }}").bold = True
p_intro.add_run(", having its registered address at {{ client.address }}, {{ client.state }}, {{ client.country }} (hereinafter referred to as \"Counterparty\").\n\n")

p_intro.add_run("(Company and Counterparty are each referred to as a \"Party\" and collectively as the \"Parties\").\n")

h2_1 = doc_mnda.add_paragraph()
h2_1.add_run("1. Purpose").bold = True
doc_mnda.add_paragraph("The Parties wish to explore and evaluate a prospective business engagement involving tech and marketing services (\"Purpose\"), in the course of which either Party may disclose non-public confidential and proprietary information to the other.")

h2_2 = doc_mnda.add_paragraph()
h2_2.add_run("2. Definition of Confidential Information").bold = True
doc_mnda.add_paragraph("\"Confidential Information\" means all technical, business, financial, operational, code, architecture, roadmap, and customer data disclosed by either Party to the other, whether orally, in writing, electronically, or by inspection.")

h2_3 = doc_mnda.add_paragraph()
h2_3.add_run("3. Obligations & Non-Disclosure").bold = True
doc_mnda.add_paragraph("Each receiving Party agrees to exercise reasonable standard of care (not less than the degree of care used for its own confidential assets) to protect against unauthorized disclosure or use. Neither Party shall disclose the Confidential Information to any third party without prior written consent.")

h2_4 = doc_mnda.add_paragraph()
h2_4.add_run("4. Term & Survival").bold = True
doc_mnda.add_paragraph("This Agreement shall remain effective for a period of {{ term_period | default('2 Years') }} from the Effective Date. The non-disclosure obligations shall survive expiration or termination for a period of {{ survival_period | default('3 Years') }}.")

h2_5 = doc_mnda.add_paragraph()
h2_5.add_run("5. Governing Law & Jurisdiction").bold = True
doc_mnda.add_paragraph("This Agreement shall be governed by the laws of India. Any dispute arising out of or related to this Agreement shall be subject to the exclusive jurisdiction of the competent courts in {{ jurisdiction | default('Kolkata, West Bengal, India') }}.")

doc_mnda.add_paragraph("\nIN WITNESS WHEREOF, the Parties have executed this Agreement by their duly authorized representatives.\n")

# Signatures Table
tbl = doc_mnda.add_table(rows=3, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.cell(0, 0).paragraphs[0].add_run("FOR: {{ company.legal_name }}").bold = True
tbl.cell(0, 1).paragraphs[0].add_run("FOR: {{ client.legal_name }}").bold = True

tbl.cell(1, 0).paragraphs[0].add_run("Authorized Signatory\nName: Authorized Signatory\nDate: {{ effective_date }}")
tbl.cell(1, 1).paragraphs[0].add_run("Authorized Signatory\nName: {{ client.signatory_name }}\nTitle: {{ client.signatory_designation }}\nDate: {{ effective_date }}")

doc_mnda.save(str(seed_dir / "mutual_nda.docx"))

# =========================================================================
# 2. ONE-WAY NDA
# =========================================================================
doc_onda = Document()
add_legal_disclaimer(doc_onda)

h1 = doc_onda.add_paragraph()
style_heading(h1, "ONE-WAY NON-DISCLOSURE AGREEMENT", size=18, bold=True)

sub = doc_onda.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = sub.add_run("Reference: {{ doc_number }}")
r_sub.font.size = Pt(9.5)
r_sub.font.color.rgb = RGBColor(100, 116, 139)

doc_onda.add_paragraph()

p_intro = doc_onda.add_paragraph()
p_intro.add_run("This Unilateral / One-Way Non-Disclosure Agreement (\"Agreement\") is entered into on ")
p_intro.add_run("{{ effective_date }}").bold = True
p_intro.add_run(" (\"Effective Date\"), by and between:\n\n")

p_intro.add_run("1. {{ company.legal_name }}").bold = True
p_intro.add_run(", having its principal office at {{ company.address }}, {{ company.state }}, India (\"Discloser\"); and\n\n")

p_intro.add_run("2. {{ client.legal_name }}").bold = True
p_intro.add_run(", having its registered office at {{ client.address }}, {{ client.state }}, {{ client.country }} (\"Recipient\").\n\n")

h2_1 = doc_onda.add_paragraph()
h2_1.add_run("1. Scope of Disclosure").bold = True
doc_onda.add_paragraph("Discloser intends to share proprietary specifications, product concepts, source code, designs, or pricing details (\"Confidential Information\") solely for the evaluation of {{ evaluation_purpose | default('commercial collaboration and technology evaluation') }}.")

h2_2 = doc_onda.add_paragraph()
h2_2.add_run("2. Duty of Confidentiality").bold = True
doc_onda.add_paragraph("Recipient shall hold in strict confidence and shall not disclose, reproduce, or commercially exploit any Confidential Information without prior written authorization from Discloser.")

h2_3 = doc_onda.add_paragraph()
h2_3.add_run("3. Term & Governing Law").bold = True
doc_onda.add_paragraph("This Agreement remains in effect for {{ term_period | default('2 Years') }}. Governing law: Laws of India, exclusive jurisdiction: {{ jurisdiction | default('Kolkata, West Bengal, India') }}.")

doc_onda.add_paragraph("\nACKNOWLEDGED AND AGREED:\n")
tbl_o = doc_onda.add_table(rows=2, cols=2)
tbl_o.cell(0, 0).paragraphs[0].add_run("DISCLOSER:\n{{ company.legal_name }}").bold = True
tbl_o.cell(0, 1).paragraphs[0].add_run("RECIPIENT:\n{{ client.legal_name }}").bold = True
tbl_o.cell(1, 0).paragraphs[0].add_run("Signatory: Authorized Representative\nDate: {{ effective_date }}")
tbl_o.cell(1, 1).paragraphs[0].add_run("Signatory: {{ client.signatory_name }}\nDesignation: {{ client.signatory_designation }}\nDate: {{ effective_date }}")

doc_onda.save(str(seed_dir / "one_way_nda.docx"))

# =========================================================================
# 3. STATEMENT OF WORK (SOW)
# =========================================================================
doc_sow = Document()
add_legal_disclaimer(doc_sow)

h1 = doc_sow.add_paragraph()
style_heading(h1, "STATEMENT OF WORK (SOW)", size=18, bold=True)

sub = doc_sow.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = sub.add_run("SOW Ref: {{ doc_number }} | Project: {{ project.name }}")
r_sub.font.size = Pt(9.5)
r_sub.font.color.rgb = RGBColor(100, 116, 139)

doc_sow.add_paragraph()

p_sow_meta = doc_sow.add_paragraph()
p_sow_meta.add_run("Effective Date: ").bold = True
p_sow_meta.add_run("{{ effective_date }}\n")
p_sow_meta.add_run("Service Provider: ").bold = True
p_sow_meta.add_run("{{ company.legal_name }} ({{ company.trade_name }})\n")
p_sow_meta.add_run("Client: ").bold = True
p_sow_meta.add_run("{{ client.legal_name }} (GSTIN: {{ client.gstin | default('N/A') }})\n")
p_sow_meta.add_run("Master Agreement / NDA Ref: ").bold = True
p_sow_meta.add_run("{{ parent_agreement_ref | default('Mutual NDA in place') }}\n")

doc_sow.add_paragraph().add_run("1. Project Scope & Deliverables").bold = True
doc_sow.add_paragraph("{{ project_scope_description }}")

doc_sow.add_paragraph().add_run("2. Milestone & Payment Schedule").bold = True
doc_sow.add_paragraph("The project fees total {{ total_fee_amount }} {{ currency | default('INR') }} (plus applicable GST). Payments are linked to milestones:")

# Milestones table with jinja loop inside docxtpl
tbl_s = doc_sow.add_table(rows=2, cols=4)
tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Milestone", "Deliverable Description", "Target Date", "Fee Amount"]
for i, h in enumerate(headers):
    tbl_s.cell(0, i).paragraphs[0].add_run(h).bold = True

tbl_s.cell(1, 0).paragraphs[0].add_run("{% for m in milestones %}{{ m.title }}")
tbl_s.cell(1, 1).paragraphs[0].add_run("{{ m.deliverable }}")
tbl_s.cell(1, 2).paragraphs[0].add_run("{{ m.due_date }}")
tbl_s.cell(1, 3).paragraphs[0].add_run("{{ m.amount }}{% if not loop.last %}\n\n{% endif %}{% endfor %}")

doc_sow.add_paragraph().add_run("3. Standard Terms & Selected Clauses").bold = True
p_c = doc_sow.add_paragraph()
p_c.add_run("{% for clause in selected_clauses %}")
p_c.add_run("• {{ clause.title }}: ").bold = True
p_c.add_run("{{ clause.body }}\n\n")
p_c.add_run("{% endfor %}")

doc_sow.add_paragraph("\nACCEPTANCE & AUTHORIZATION:\n")
tbl_sow_sig = doc_sow.add_table(rows=2, cols=2)
tbl_sow_sig.cell(0, 0).paragraphs[0].add_run("FOR: {{ company.legal_name }}").bold = True
tbl_sow_sig.cell(0, 1).paragraphs[0].add_run("FOR: {{ client.legal_name }}").bold = True
tbl_sow_sig.cell(1, 0).paragraphs[0].add_run("Authorized Signatory\nDate: {{ effective_date }}")
tbl_sow_sig.cell(1, 1).paragraphs[0].add_run("Signatory: {{ client.signatory_name }}\nDesignation: {{ client.signatory_designation }}\nDate: {{ effective_date }}")

doc_sow.save(str(seed_dir / "statement_of_work.docx"))

print("Created seed templates: mutual_nda.docx, one_way_nda.docx, statement_of_work.docx")
